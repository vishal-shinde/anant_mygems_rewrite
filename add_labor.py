import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
import psycopg2
from psycopg2 import Error

# ============================================================
# DATABASE CONFIGURATION
# ============================================================
DB_CONFIG = {
    'host': 'localhost',
    'database': 'mygems',
    'user': 'myuser',
    'password': '28116',
    'port': '5432'
}

class LabourManagementApp:
    def __init__(self, root):
        self.root = root
        self.root.title("🏭 Labour Management System - MyGems")
        self.root.geometry("1280x720")
        self.root.minsize(1100, 650)
        
        self.colors = {
            'primary': '#6366f1', 'secondary': '#10b981', 'accent': '#f43f5e',
            'warning': '#f59e0b', 'info': '#06b6d4', 'dark': '#1e293b',
            'light': '#f8fafc', 'white': '#ffffff', 'gray': '#64748b',
            'success': '#22c55e', 'purple': '#a855f7', 'teal': '#14b8a6',
            'orange': '#fb923c', 'danger': '#ef4444'
        }
        
        self.entries = []
        self.labour_summary = {}
        self.workers_dict = {}
        self.workers_with_subtype = {}
        self.editing_index = None  # Track which entry is being edited
        
        self.db_connection = None
        self.current_txn_id = None
        self.daily_counter = 1
        
        self.setup_styles()
        self.setup_ui()
        
        if self.connect_database():
            self.load_workers()
            self.load_processes()
            self.get_daily_counter()
            self.generate_txn_id()
            self.update_worker_info()
    
    def setup_styles(self):
        style = ttk.Style()
        try:
            style.theme_use('clam')
        except:
            pass
        
        style.configure("Custom.Treeview",
                        background=self.colors['white'],
                        foreground=self.colors['dark'],
                        fieldbackground=self.colors['white'],
                        rowheight=28,
                        font=('Segoe UI', 10))
        
        style.configure("Custom.Treeview.Heading",
                        background=self.colors['primary'],
                        foreground=self.colors['white'],
                        font=('Segoe UI', 10, 'bold'))
        
        style.map("Custom.Treeview",
                  background=[('selected', self.colors['info'])])
        
        style.configure("Custom.TCombobox",
                        fieldbackground=self.colors['white'],
                        font=('Segoe UI', 10))
    
    # ============================================================
    # DATABASE FUNCTIONS
    # ============================================================
    def connect_database(self):
        try:
            self.db_connection = psycopg2.connect(**DB_CONFIG)
            self.db_connection.autocommit = False
            self.status_label.config(text="● Connected", fg=self.colors['success'])
            return True
        except Error as e:
            messagebox.showerror("Database Error", f"❌ Connection failed:\n{str(e)}")
            return False
    
    def load_workers(self):
        try:
            cursor = self.db_connection.cursor()
            
            # We filter by type_id = 5 based on your screenshot (WORKER)
            query = """
                SELECT account_id, account_name 
                FROM accounts 
                WHERE type_id = 5
                ORDER BY account_name
            """
            cursor.execute(query)
            workers = cursor.fetchall()
            cursor.close()
            
            if workers:
                self.workers_dict = {}
                worker_list = []
                
                for account_id, name in workers:
                    self.workers_dict[name] = account_id
                    worker_list.append(name)
                
                self.worker_combo['values'] = worker_list
                if worker_list:
                    self.worker_combo.current(0)
        except Error as e:
            messagebox.showerror("Database Error", f"❌ Failed to load workers:\n{str(e)}")
    
    def load_processes(self):
        processes = ['Wax', 'filing', 'Laser', 'Setting', 
                     'Polish', 'QC', 'Mold-Making', 'Other']
        self.process_combo['values'] = processes
        if processes:
            self.process_combo.current(0)
    
    def get_daily_counter(self):
        try:
            cursor = self.db_connection.cursor()
            query = """SELECT COUNT(*) FROM labor WHERE date = %s"""
            cursor.execute(query, (datetime.now().strftime("%Y-%m-%d"),))
            count = cursor.fetchone()[0]
            cursor.close()
            self.daily_counter = count + 1
        except Error:
            self.daily_counter = 1
    
    def generate_txn_id(self):
        date_part = datetime.now().strftime("%d%m%y")
        self.current_txn_id = f"LAB-{date_part}{self.daily_counter:03d}"
        self.txn_id_entry.delete(0, 'end')
        self.txn_id_entry.insert(0, self.current_txn_id)
    
    def verify_lot_numbers(self, lot_numbers):
        """Verify lot numbers exist in gold_jewelry or silver_jewelry tables.
        Returns set of valid lot numbers."""
        valid_lots = set()
        try:
            cursor = self.db_connection.cursor()
            
            # Check gold_jewelry
            query_gold = "SELECT lot_no FROM gold_jewelry WHERE lot_no = ANY(%s)"
            cursor.execute(query_gold, (list(lot_numbers),))
            for row in cursor.fetchall():
                valid_lots.add(str(row[0]))
            
            # Check silver_jewelry
            query_silver = "SELECT lot_no FROM silver_jewelry WHERE lot_no = ANY(%s)"
            cursor.execute(query_silver, (list(lot_numbers),))
            for row in cursor.fetchall():
                valid_lots.add(str(row[0]))
            
            cursor.close()
        except Error as e:
            messagebox.showerror("Database Error", f"❌ Failed to verify lots:\n{str(e)}")
        
        return valid_lots
    
    def save_labour_to_db(self):
        if not self.entries:
            return False, []
        
        try:
            cursor = self.db_connection.cursor()
            
            process = self.process_combo.get()
            labour_date = self.date_entry.get()
            txn_id = self.txn_id_entry.get() # This is your 'LAB-...' string
            
            # --- UPDATED QUERY ---
            # Added 'txn_id' to the INSERT columns
            insert_query = """
                INSERT INTO labor 
                (date, lot_no, worker_id, process, labor_amount, status, remarks, txn_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING job_no
            """
            
            saved_jobs = []
            remarks = "Batch Entry" # Or leave empty if you prefer
            
            for entry in self.entries:
                worker_id = entry.get('worker_id')
                
                cursor.execute(insert_query, (
                    labour_date,
                    entry['lot_no'],
                    worker_id,
                    entry['process'],
                    entry['amount'],
                    'PENDING',
                    remarks,
                    txn_id # Saving the string 'LAB-28116...' here
                ))
                job_no = cursor.fetchone()[0]
                saved_jobs.append(job_no)
            
            self.db_connection.commit()
            cursor.close()
            return True, saved_jobs
            
        except Error as e:
            self.db_connection.rollback()
            messagebox.showerror("Database Error", f"❌ Failed to save entry:\n\n{str(e)}")
            return False, []
    
    # ============================================================
    # UI SETUP
    # ============================================================
    def setup_ui(self):
        # HEADER
        header_frame = tk.Frame(self.root, bg=self.colors['primary'], height=70)
        header_frame.pack(fill='x')
        header_frame.pack_propagate(False)
        
        title_frame = tk.Frame(header_frame, bg=self.colors['primary'])
        title_frame.pack(side='left', padx=20, pady=10)
        
        tk.Label(title_frame, text="🏭", font=('Segoe UI Emoji', 28), 
                 bg=self.colors['primary'], fg=self.colors['white']).pack(side='left')
        tk.Label(title_frame, text="Add Labour", 
                 font=('Segoe UI', 24, 'bold'), 
                 fg=self.colors['white'], bg=self.colors['primary']).pack(side='left', padx=8)
        tk.Label(title_frame, text="| MyGems", 
                 font=('Segoe UI', 12), 
                 fg='#cbd5e1', bg=self.colors['primary']).pack(side='left')
        
        self.status_label = tk.Label(header_frame, text="● Connecting...", 
                                      font=('Segoe UI', 10, 'bold'),
                                      fg=self.colors['warning'], bg=self.colors['primary'])
        self.status_label.pack(side='right', padx=15)
        
        main_container = tk.Frame(self.root, bg=self.colors['light'])
        main_container.pack(fill='both', expand=True, padx=12, pady=8)
        
        # TRANSACTION INFO
        txn_frame = tk.Frame(main_container, bg=self.colors['white'], relief='solid', bd=1)
        txn_frame.pack(fill='x', pady=(0, 8))
        
        date_container = tk.Frame(txn_frame, bg=self.colors['white'])
        date_container.pack(side='left', padx=20, pady=10)
        
        tk.Label(date_container, text="📅 Date:", font=('Segoe UI', 11, 'bold'), 
                 bg=self.colors['white'], fg=self.colors['dark']).pack(side='left')
        self.date_entry = tk.Entry(date_container, font=('Segoe UI', 11), 
                                    width=12, bd=2, relief='groove')
        self.date_entry.pack(side='left', padx=5)
        self.date_entry.insert(0, datetime.now().strftime("%Y-%m-%d"))
        
        txn_container = tk.Frame(txn_frame, bg=self.colors['white'])
        txn_container.pack(side='left', padx=20, pady=10)
        
        tk.Label(txn_container, text="🆔 Txn ID:", font=('Segoe UI', 11, 'bold'), 
                 bg=self.colors['white'], fg=self.colors['dark']).pack(side='left')
        self.txn_id_entry = tk.Entry(txn_container, font=('Segoe UI', 11), 
                                      width=18, bd=2, relief='groove')
        self.txn_id_entry.pack(side='left', padx=5)
        
        # WORKER & PROCESS
        selection_frame = tk.Frame(main_container, bg=self.colors['white'], relief='solid', bd=1)
        selection_frame.pack(fill='x', pady=(0, 8))
        
        worker_container = tk.Frame(selection_frame, bg=self.colors['white'])
        worker_container.pack(side='left', padx=20, pady=10)
        
        tk.Label(worker_container, text="👷 Worker:", font=('Segoe UI', 11, 'bold'), 
                 bg=self.colors['white'], fg=self.colors['dark']).pack(side='left')
        self.worker_combo = ttk.Combobox(worker_container, font=('Segoe UI', 10),
                                          state='readonly', width=22,
                                          style="Custom.TCombobox")
        self.worker_combo.pack(side='left', padx=5)
        self.worker_combo.bind('<<ComboboxSelected>>', lambda e: self.update_worker_info())
        
        process_container = tk.Frame(selection_frame, bg=self.colors['white'])
        process_container.pack(side='left', padx=20, pady=10)
        
        tk.Label(process_container, text="⚙️ Process:", font=('Segoe UI', 11, 'bold'), 
                 bg=self.colors['white'], fg=self.colors['dark']).pack(side='left')
        self.process_combo = ttk.Combobox(process_container, font=('Segoe UI', 10),
                                           state='readonly', width=18,
                                           style="Custom.TCombobox")
        self.process_combo.pack(side='left', padx=5)
        
        # DATA ENTRY
        entry_section = tk.Frame(main_container, bg=self.colors['white'], relief='solid', bd=1)
        entry_section.pack(fill='x', pady=(0, 8))
        
        entry_inner = tk.Frame(entry_section, bg=self.colors['white'])
        entry_inner.pack(pady=10)
        
        tk.Label(entry_inner, text="📝", font=('Segoe UI Emoji', 14), 
                 bg=self.colors['white']).pack(side='left', padx=(10, 5))
        
        tk.Label(entry_inner, text="Lot No:", font=('Segoe UI', 11, 'bold'), 
                 bg=self.colors['white'], fg=self.colors['gray']).pack(side='left')
        self.lot_entry = tk.Entry(entry_inner, font=('Segoe UI', 11), 
                                   width=15, bd=2, relief='groove')
        self.lot_entry.pack(side='left', padx=5)
        
        tk.Label(entry_inner, text="Amount (₹):", font=('Segoe UI', 11, 'bold'), 
                 bg=self.colors['white'], fg=self.colors['gray']).pack(side='left', padx=(15, 0))
        self.amount_entry = tk.Entry(entry_inner, font=('Segoe UI', 11), 
                                      width=15, bd=2, relief='groove')
        self.amount_entry.pack(side='left', padx=5)
        
        self.add_btn = tk.Button(entry_inner, text="➕ Add to List", 
                            font=('Segoe UI', 10, 'bold'),
                            bg=self.colors['secondary'], fg=self.colors['white'],
                            activebackground=self.colors['success'],
                            padx=15, pady=6, bd=0, cursor='hand2',
                            command=self.add_entry)
        self.add_btn.pack(side='left', padx=(20, 5))
        
        # IMPORT button beside Add to List
        import_btn = tk.Button(entry_inner, text="📥 Import Excel", 
                            font=('Segoe UI', 10, 'bold'),
                            bg=self.colors['teal'], fg=self.colors['white'],
                            activebackground=self.colors['info'],
                            padx=15, pady=6, bd=0, cursor='hand2',
                            command=self.import_excel)
        import_btn.pack(side='left', padx=5)
        
        self.lot_entry.bind('<Return>', lambda e: self.amount_entry.focus())
        self.amount_entry.bind('<Return>', lambda e: self.add_entry())
        
        # DISPLAY SECTION
        display_container = tk.Frame(main_container, bg=self.colors['light'])
        display_container.pack(fill='both', expand=True, pady=(0, 8))
        
        # LEFT - Excel Sheet
        left_frame = tk.Frame(display_container, bg=self.colors['white'], relief='solid', bd=1)
        left_frame.pack(side='left', fill='both', expand=True, padx=(0, 6))
        
        left_header = tk.Frame(left_frame, bg=self.colors['info'])
        left_header.pack(fill='x')
        tk.Label(left_header, text="📊 Labour Entries", 
                 font=('Segoe UI', 12, 'bold'),
                 fg=self.colors['white'], bg=self.colors['info']).pack(pady=6, padx=12)
        
        tree_container = tk.Frame(left_frame, bg=self.colors['white'])
        tree_container.pack(fill='both', expand=True, padx=8, pady=8)
        
        y_scroll = ttk.Scrollbar(tree_container, orient='vertical')
        y_scroll.pack(side='right', fill='y')
        
        columns = ('Lot No', 'Amount', 'Worker', 'Process')
        self.tree = ttk.Treeview(tree_container, columns=columns, show='headings',
                                  yscrollcommand=y_scroll.set,
                                  style="Custom.Treeview", height=10)
        
        y_scroll.config(command=self.tree.yview)
        
        self.tree.heading('Lot No', text='📦 Lot No')
        self.tree.heading('Amount', text='💰 Amount')
        self.tree.heading('Worker', text='👷 Worker')
        self.tree.heading('Process', text='⚙️ Process')
        
        self.tree.column('Lot No', width=110, anchor='center')
        self.tree.column('Amount', width=110, anchor='center')
        self.tree.column('Worker', width=160, anchor='center')
        self.tree.column('Process', width=130, anchor='center')
        
        self.tree.pack(fill='both', expand=True)
        
        self.tree.tag_configure('even', background='#f1f5f9')
        self.tree.tag_configure('odd', background='#ffffff')
        self.tree.tag_configure('invalid', background='#fecaca', foreground='#991b1b')
        self.tree.tag_configure('editing', background='#fde68a', foreground='#92400e')
        
        # Double-click to edit
        self.tree.bind('<Double-1>', lambda e: self.edit_entry())
        
        # RIGHT - Summary Dashboard
        right_frame = tk.Frame(display_container, bg=self.colors['white'], 
                               relief='solid', bd=1, width=380)
        right_frame.pack(side='right', fill='both', padx=(6, 0))
        right_frame.pack_propagate(False)
        
        right_header = tk.Frame(right_frame, bg=self.colors['accent'])
        right_header.pack(fill='x')
        tk.Label(right_header, text="📈 Summary Dashboard", 
                 font=('Segoe UI', 12, 'bold'),
                 fg=self.colors['white'], bg=self.colors['accent']).pack(pady=6, padx=12)
        
        session_frame = tk.Frame(right_frame, bg=self.colors['light'], relief='solid', bd=1)
        session_frame.pack(fill='x', padx=8, pady=8)
        
        tk.Label(session_frame, text="📋 Current Session", 
                 font=('Segoe UI', 10, 'bold'),
                 bg=self.colors['light'], fg=self.colors['primary']).pack(anchor='w', padx=8, pady=4)
        
        session_items = [
            ('👷 Worker:', 'worker_info', self.colors['dark']),
            ('⚙️ Process:', 'process_info', self.colors['info']),
            ('🆔 Txn ID:', 'txn_info', self.colors['gray']),
        ]
        
        for label, key, color in session_items:
            item = tk.Frame(session_frame, bg=self.colors['light'])
            item.pack(fill='x', padx=8, pady=1)
            tk.Label(item, text=label, font=('Segoe UI', 9, 'bold'), 
                     bg=self.colors['light'], fg=self.colors['gray'], 
                     width=10, anchor='w').pack(side='left')
            value_label = tk.Label(item, text="-", font=('Segoe UI', 9), 
                                    bg=self.colors['light'], fg=color, anchor='w')
            value_label.pack(side='left', fill='x', expand=True)
            setattr(self, f'{key}_label', value_label)
        
        tk.Frame(right_frame, bg='#e2e8f0', height=1).pack(fill='x', padx=8, pady=4)
        
        labour_summary_header = tk.Frame(right_frame, bg=self.colors['purple'])
        labour_summary_header.pack(fill='x', padx=8, pady=(0, 4))
        tk.Label(labour_summary_header, text="👥 Labour-wise Summary", 
                 font=('Segoe UI', 11, 'bold'),
                 fg=self.colors['white'], bg=self.colors['purple']).pack(pady=4, padx=8)
        
        summary_tree_frame = tk.Frame(right_frame, bg=self.colors['white'])
        summary_tree_frame.pack(fill='both', expand=True, padx=8, pady=4)
        
        summary_scroll = ttk.Scrollbar(summary_tree_frame, orient='vertical')
        summary_scroll.pack(side='right', fill='y')
        
        self.summary_tree = ttk.Treeview(summary_tree_frame, 
                                          columns=('Worker', 'Items', 'Amount'),
                                          show='headings',
                                          yscrollcommand=summary_scroll.set,
                                          style="Custom.Treeview", height=6)
        
        summary_scroll.config(command=self.summary_tree.yview)
        
        self.summary_tree.heading('Worker', text='👷 Worker')
        self.summary_tree.heading('Items', text='Items')
        self.summary_tree.heading('Amount', text='Amount (₹)')
        
        self.summary_tree.column('Worker', width=140, anchor='w')
        self.summary_tree.column('Items', width=50, anchor='center')
        self.summary_tree.column('Amount', width=110, anchor='e')
        
        self.summary_tree.pack(fill='both', expand=True)
        
        self.summary_tree.tag_configure('even', background='#f1f5f9')
        self.summary_tree.tag_configure('odd', background='#ffffff')
        
        # GRAND TOTAL
        grand_total_frame = tk.Frame(right_frame, bg=self.colors['secondary'], relief='raised', bd=2)
        grand_total_frame.pack(fill='x', padx=8, pady=6)
        
        tk.Label(grand_total_frame, text="💰 GRAND TOTAL", 
                 font=('Segoe UI', 10, 'bold'),
                 bg=self.colors['secondary'], fg=self.colors['white']).pack(anchor='w', padx=8, pady=3)
        
        total_row = tk.Frame(grand_total_frame, bg=self.colors['secondary'])
        total_row.pack(fill='x', padx=8, pady=(0, 6))
        
        self.grand_items_label = tk.Label(total_row, text="0 items", 
                                           font=('Segoe UI', 10),
                                           bg=self.colors['secondary'], fg=self.colors['white'])
        self.grand_items_label.pack(side='left')
        
        self.grand_amount_label = tk.Label(total_row, text="₹ 0.00", 
                                            font=('Segoe UI', 14, 'bold'),
                                            bg=self.colors['secondary'], fg=self.colors['white'])
        self.grand_amount_label.pack(side='right')
        
        # BOTTOM BUTTONS - With EDIT and DELETE
        button_frame = tk.Frame(self.root, bg=self.colors['dark'], height=60)
        button_frame.pack(fill='x', side='bottom')
        button_frame.pack_propagate(False)
        
        btn_container = tk.Frame(button_frame, bg=self.colors['dark'])
        btn_container.pack(expand=True)
        
        buttons_config = [
            ('💾 Save to DB', self.save_data, self.colors['secondary'], self.colors['success']),
            ('✨ New Entry', self.new_entry, self.colors['info'], self.colors['primary']),
            ('✏️ Edit', self.edit_entry, self.colors['warning'], self.colors['orange']),
            ('🗑️ Delete', self.delete_entry, self.colors['danger'], self.colors['accent']),
            ('🔄 Refresh', self.refresh_data, self.colors['purple'], self.colors['primary']),
            ('❌ Exit', self.exit_app, self.colors['accent'], self.colors['warning']),
        ]
        
        for text, command, bg, active_bg in buttons_config:
            btn = tk.Button(btn_container, text=text,
                            font=('Segoe UI', 11, 'bold'),
                            bg=bg, fg=self.colors['white'],
                            activebackground=active_bg,
                            padx=18, pady=8,
                            bd=0, cursor='hand2',
                            command=command)
            btn.pack(side='left', padx=8, pady=12)
    
    # ============================================================
    # EVENT HANDLERS
    # ============================================================
    def update_worker_info(self):
        selected = self.worker_combo.get()
        if selected:
            self.worker_info_label.config(text=selected)
        else:
            self.worker_info_label.config(text="-")
    
    def add_entry(self):
        lot_no = self.lot_entry.get().strip()
        amount_text = self.amount_entry.get().strip()
        worker = self.worker_combo.get()
        process = self.process_combo.get()
        
        if not lot_no:
            messagebox.showwarning("Warning", "Please enter Lot No!")
            self.lot_entry.focus()
            return
        
        if not amount_text:
            messagebox.showwarning("Warning", "Please enter Amount!")
            self.amount_entry.focus()
            return
        
        try:
            amount = float(amount_text)
        except ValueError:
            messagebox.showerror("Error", "Please enter a valid numeric amount!")
            self.amount_entry.focus()
            return
        
        if not worker:
            messagebox.showwarning("Warning", "Please select a worker!")
            return
        
        if not process:
            messagebox.showwarning("Warning", "Please select a process!")
            return
        
        # Check if we're editing an existing entry
        if self.editing_index is not None:
            # Update existing entry
            old_entry = self.entries[self.editing_index]
            
            # Update labour summary - remove old
            old_worker = old_entry['worker']
            self.labour_summary[old_worker]['items'] -= 1
            self.labour_summary[old_worker]['amount'] -= old_entry['amount']
            if self.labour_summary[old_worker]['items'] <= 0:
                del self.labour_summary[old_worker]
            
            # Update entry
            self.entries[self.editing_index] = {
                'lot_no': lot_no,
                'amount': amount,
                'worker': worker,
                'worker_id': self.workers_dict.get(worker),
                'process': process
            }
            
            # Add to summary - new
            if worker not in self.labour_summary:
                self.labour_summary[worker] = {'items': 0, 'amount': 0.0}
            self.labour_summary[worker]['items'] += 1
            self.labour_summary[worker]['amount'] += amount
            
            self.editing_index = None
            self.add_btn.config(text="➕ Add to List", bg=self.colors['secondary'])
            self.refresh_tree()
        else:
            # Add new entry
            self.entries.append({
                'lot_no': lot_no, 
                'amount': amount,
                'worker': worker,
                'worker_id': self.workers_dict.get(worker),
                'process': process
            })
            
            if worker not in self.labour_summary:
                self.labour_summary[worker] = {'items': 0, 'amount': 0.0}
            self.labour_summary[worker]['items'] += 1
            self.labour_summary[worker]['amount'] += amount
            
            self.refresh_tree()
        
        # Clear fields
        self.lot_entry.delete(0, 'end')
        self.amount_entry.delete(0, 'end')
        self.lot_entry.focus()
        
        self.update_summary()
    
    def refresh_tree(self, invalid_lots=None):
        """Refresh treeview from self.entries list"""
        if invalid_lots is None:
            invalid_lots = set()
        
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        for idx, entry in enumerate(self.entries):
            if str(entry['lot_no']) in invalid_lots:
                tag = 'invalid'
            else:
                tag = 'even' if idx % 2 == 0 else 'odd'
            
            self.tree.insert('', 'end',
                            iid=str(idx),
                            values=(entry['lot_no'], 
                                    f"₹ {entry['amount']:,.2f}",
                                    entry['worker'],
                                    entry['process']),
                            tags=(tag,))
    
    def edit_entry(self):
        """Edit selected entry"""
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Warning", "Please select an entry to edit!")
            return
        
        idx = int(selected[0])
        if idx < 0 or idx >= len(self.entries):
            return
        
        entry = self.entries[idx]
        
        # Populate fields
        self.lot_entry.delete(0, 'end')
        self.lot_entry.insert(0, entry['lot_no'])
        
        self.amount_entry.delete(0, 'end')
        self.amount_entry.insert(0, str(entry['amount']))
        
        if entry['worker'] in self.worker_combo['values']:
            self.worker_combo.set(entry['worker'])
            self.update_worker_info()
        
        if entry['process'] in self.process_combo['values']:
            self.process_combo.set(entry['process'])
        
        self.editing_index = idx
        self.add_btn.config(text="✅ Update Entry", bg=self.colors['warning'])
        
        # Highlight in tree
        self.refresh_tree()
        try:
            self.tree.item(str(idx), tags=('editing',))
            self.tree.selection_set(str(idx))
        except:
            pass
        
        self.lot_entry.focus()
    
    def delete_entry(self):
        """Delete selected entry/entries"""
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Warning", "Please select entry/entries to delete!")
            return
        
        confirm = messagebox.askyesno("Confirm Delete", 
            f"Delete {len(selected)} selected entry/entries?")
        if not confirm:
            return
        
        # Get indices and sort descending so deletion doesn't shift indices
        indices = sorted([int(sel) for sel in selected], reverse=True)
        
        for idx in indices:
            if 0 <= idx < len(self.entries):
                entry = self.entries[idx]
                worker = entry['worker']
                
                # Update labour summary
                if worker in self.labour_summary:
                    self.labour_summary[worker]['items'] -= 1
                    self.labour_summary[worker]['amount'] -= entry['amount']
                    if self.labour_summary[worker]['items'] <= 0:
                        del self.labour_summary[worker]
                
                del self.entries[idx]
        
        # Reset editing state if editing was deleted
        if self.editing_index is not None and self.editing_index in indices:
            self.editing_index = None
            self.add_btn.config(text="➕ Add to List", bg=self.colors['secondary'])
            self.lot_entry.delete(0, 'end')
            self.amount_entry.delete(0, 'end')
        
        self.refresh_tree()
        self.update_summary()
    
    def import_excel(self):
        """Import lot numbers and amounts from Excel file"""
        worker = self.worker_combo.get()
        process = self.process_combo.get()
        
        if not worker:
            messagebox.showwarning("Warning", "Please select a worker before importing!")
            return
        
        if not process:
            messagebox.showwarning("Warning", "Please select a process before importing!")
            return
        
        file_path = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
        
        if not file_path:
            return
        
        try:
            # Try openpyxl first
            try:
                from openpyxl import load_workbook
                wb = load_workbook(file_path, data_only=True)
                ws = wb.active
                rows_data = list(ws.iter_rows(values_only=True))
            except ImportError:
                # Fallback to pandas
                try:
                    import pandas as pd
                    df = pd.read_excel(file_path, header=None)
                    rows_data = df.values.tolist()
                except ImportError:
                    messagebox.showerror("Missing Library", 
                        "Please install openpyxl or pandas:\npip install openpyxl")
                    return
            
            if not rows_data:
                messagebox.showwarning("Empty File", "The Excel file is empty!")
                return
            
            # Detect if first row is header
            first_row = rows_data[0]
            start_idx = 0
            if first_row and len(first_row) >= 2:
                first_val = str(first_row[0]).strip().lower() if first_row[0] is not None else ""
                second_val = str(first_row[1]).strip().lower() if first_row[1] is not None else ""
                if 'lot' in first_val or 'amount' in second_val or 'amt' in second_val:
                    start_idx = 1
            
            imported_count = 0
            error_rows = []
            
            for row_idx, row in enumerate(rows_data[start_idx:], start=start_idx + 1):
                if not row or len(row) < 2:
                    continue
                
                lot_no = row[0]
                amount = row[1]
                
                if lot_no is None or amount is None:
                    continue
                
                lot_no = str(lot_no).strip()
                if not lot_no:
                    continue
                
                try:
                    amount = float(amount)
                except (ValueError, TypeError):
                    error_rows.append(row_idx)
                    continue
                
                # Add to entries
                self.entries.append({
                    'lot_no': lot_no,
                    'amount': amount,
                    'worker': worker,
                    'worker_id': self.workers_dict.get(worker),
                    'process': process
                })
                
                if worker not in self.labour_summary:
                    self.labour_summary[worker] = {'items': 0, 'amount': 0.0}
                self.labour_summary[worker]['items'] += 1
                self.labour_summary[worker]['amount'] += amount
                
                imported_count += 1
            
            self.refresh_tree()
            self.update_summary()
            
            msg = f"✅ Successfully imported {imported_count} entries!"
            if error_rows:
                msg += f"\n\n⚠️ Skipped rows (invalid data): {error_rows[:10]}"
            messagebox.showinfo("Import Complete", msg)
            
        except Exception as e:
            messagebox.showerror("Import Error", f"❌ Failed to import:\n{str(e)}")
    
    def update_summary(self):
        self.process_info_label.config(text=self.process_combo.get() or "-")
        txn_id = self.txn_id_entry.get() or "-"
        self.txn_info_label.config(text=txn_id)
        
        for item in self.summary_tree.get_children():
            self.summary_tree.delete(item)
        
        total_items = 0
        total_amount = 0.0
        
        for idx, (worker, data) in enumerate(self.labour_summary.items()):
            tag = 'even' if idx % 2 == 0 else 'odd'
            self.summary_tree.insert('', 'end',
                                      values=(worker, data['items'], f"₹ {data['amount']:,.2f}"),
                                      tags=(tag,))
            total_items += data['items']
            total_amount += data['amount']
        
        self.grand_items_label.config(text=f"🔢 {total_items} items")
        self.grand_amount_label.config(text=f"₹ {total_amount:,.2f}")
    
    def save_data(self):
        if not self.entries:
            messagebox.showwarning("Warning", "No entries to save!")
            return
        
        if not self.db_connection:
            messagebox.showerror("Error", "Database not connected!")
            return
        
        # ====== VERIFY LOT NUMBERS ======
        all_lots = list(set(str(e['lot_no']) for e in self.entries))
        valid_lots = self.verify_lot_numbers(all_lots)
        
        invalid_lots = set(all_lots) - valid_lots
        
        if invalid_lots:
            # Highlight invalid lots in display
            self.refresh_tree(invalid_lots=invalid_lots)
            
            invalid_list = sorted(invalid_lots)
            preview = ', '.join(invalid_list[:10])
            if len(invalid_list) > 10:
                preview += f" ... (+{len(invalid_list) - 10} more)"
            
            messagebox.showerror(
                "❌ Invalid Lot Numbers",
                f"The following lot numbers do not exist in gold_jewelry or silver_jewelry tables:\n\n"
                f"{preview}\n\n"
                f"Total invalid: {len(invalid_lots)}\n\n"
                f"Highlighted in red. Please correct or remove them before saving."
            )
            return
        
        # All lots valid - proceed with save
        summary_text = f"📋 Save Summary:\n\n"
        summary_text += f"   ⚙️ Process: {self.process_combo.get()}\n"
        summary_text += f"   🔢 Total Items: {sum(d['items'] for d in self.labour_summary.values())}\n"
        summary_text += f"   💰 Total Amount: ₹ {sum(d['amount'] for d in self.labour_summary.values()):,.2f}\n"
        summary_text += f"   🆔 Txn ID: {self.txn_id_entry.get()}\n"
        summary_text += f"   ✅ All {len(all_lots)} lots verified\n\n"
        summary_text += f"   Save to database?"
        
        confirm = messagebox.askyesno("Confirm Save", summary_text)
        
        if confirm:
            result = self.save_labour_to_db()
            if result:
                success, job_nos = result
                if success:
                    self.daily_counter += 1
                    messagebox.showinfo("✅ Success", 
                                       f"Data saved successfully!\n\n"
                                       f"   Job Numbers: {', '.join(map(str, job_nos[:5]))}\n"
                                       f"   Total: {len(job_nos)} entries")
                    self.new_entry()
    
    def new_entry(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        self.entries = []
        self.labour_summary = {}
        self.editing_index = None
        self.add_btn.config(text="➕ Add to List", bg=self.colors['secondary'])
        
        self.lot_entry.delete(0, 'end')
        self.amount_entry.delete(0, 'end')
        
        self.get_daily_counter()
        self.generate_txn_id()
        self.update_summary()
        self.update_worker_info()
        self.lot_entry.focus()
    
    def refresh_data(self):
        if self.db_connection:
            self.load_workers()
            messagebox.showinfo("Refreshed", "✅ Workers refreshed!")
        else:
            messagebox.showwarning("Warning", "Database not connected!")
    
    def exit_app(self):
        if self.entries:
            confirm = messagebox.askyesnocancel("Exit", 
                "⚠️ You have unsaved entries!\nSave before exiting?")
            if confirm is None:
                return
            elif confirm:
                self.save_data()
                return
        
        if self.db_connection:
            self.db_connection.close()
        self.root.destroy()


def main():
    root = tk.Tk()
    app = LabourManagementApp(root)
    root.protocol("WM_DELETE_WINDOW", app.exit_app)
    root.mainloop()

if __name__ == "__main__":
    main()
